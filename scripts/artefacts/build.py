"""Generate CloudMigrate's downloadable artefacts for every cloud.

Usage:
    python3 -m venv .venv && .venv/bin/pip install -r scripts/artefacts/requirements.txt
    .venv/bin/python scripts/artefacts/build.py

Writes files to public/downloads/<cloud>/ and a manifest to src/data/artefacts.json.
"""

from __future__ import annotations

import datetime as dt
import json
import os
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Cm
from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, os.path.dirname(__file__))
from profiles import PROFILES  # noqa: E402

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
OUT = os.path.join(ROOT, "public", "downloads")
MANIFEST = os.path.join(ROOT, "src", "data", "artefacts.json")
VERSION = "v1.0"
TODAY = dt.date.today().strftime("%d %b %Y")

INK = "15191C"
PAPER = "F3F0E8"
RULE = "D4CEBF"
ACCENT = "C8442B"

thin = Side(style="thin", color=RULE)
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
HEAD_FILL = PatternFill("solid", fgColor=INK)
HEAD_FONT = Font(name="Calibri", bold=True, color="FFFFFF", size=10)
BODY_FONT = Font(name="Calibri", size=10, color=INK)
WRAP = Alignment(wrap_text=True, vertical="top")
STATUS = ["Not started", "In progress", "Done", "Blocked", "N/A"]
PRIORITY = ["Critical", "High", "Medium", "Low"]
RACI_ROLES = ["Customer IT lead", "Application owner", "CloudMigrate architect", "Migration engineer", "Security lead", "Network lead", "DBA", "Service desk"]


# ---------------------------------------------------------------- xlsx helpers

def cover_sheet(wb: Workbook, title: str, p: dict, purpose: str, how: list[str]):
    ws = wb.active
    ws.title = "Cover"
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 28
    ws.column_dimensions["C"].width = 90
    ws["B2"] = "CloudMigrate.in"
    ws["B2"].font = Font(name="Calibri", size=11, bold=True, color=ACCENT)
    ws["B3"] = title
    ws["B3"].font = Font(name="Calibri", size=20, bold=True, color=INK)
    ws["B4"] = f"{p['name']} · {VERSION} · generated {TODAY}"
    ws["B4"].font = Font(name="Calibri", size=10, color="676C6E")
    ws.row_dimensions[3].height = 30
    row = 6
    for k, v in [("Purpose", purpose), ("Customer", ""), ("Project / programme", ""), ("Prepared by", ""), ("Reviewed by", ""), ("Last updated", "")]:
        ws.cell(row=row, column=2, value=k).font = Font(name="Calibri", bold=True, size=10, color=INK)
        c = ws.cell(row=row, column=3, value=v)
        c.font = BODY_FONT
        c.alignment = WRAP
        c.border = Border(bottom=thin)
        row += 1
    row += 1
    ws.cell(row=row, column=2, value="How to use").font = Font(name="Calibri", bold=True, size=10, color=INK)
    for i, line in enumerate(how):
        c = ws.cell(row=row + i, column=3, value=f"{i + 1}. {line}")
        c.font = BODY_FONT
        c.alignment = WRAP
    row += len(how) + 1
    ws.cell(row=row, column=2, value="Reference").font = Font(name="Calibri", bold=True, size=10, color=INK)
    refs = [
        ("Target regions", p["regions"]),
        ("Landing zone", p["landing_zone"]),
        ("Security baseline", p["benchmark"]),
    ]
    for i, (k, v) in enumerate(refs):
        ws.cell(row=row + i, column=3, value=f"{k}: {v}").font = BODY_FONT
    ws.cell(row=row + len(refs) + 2, column=2, value="Template provided as-is for planning purposes. Validate vendor programme terms and service availability before relying on them.").font = Font(name="Calibri", italic=True, size=9, color="676C6E")
    return ws


def table_sheet(wb: Workbook, name: str, headers: list[tuple[str, int]], rows: list[list], status_col: str | None = None, priority_col: str | None = None, freeze="C2"):
    ws = wb.create_sheet(name)
    ws.sheet_view.showGridLines = False
    for i, (h, w) in enumerate(headers, start=1):
        c = ws.cell(row=1, column=i, value=h)
        c.fill = HEAD_FILL
        c.font = HEAD_FONT
        c.alignment = Alignment(wrap_text=True, vertical="center")
        c.border = BORDER
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[1].height = 30
    for r, row in enumerate(rows, start=2):
        for ci, val in enumerate(row, start=1):
            c = ws.cell(row=r, column=ci, value=val)
            c.font = BODY_FONT
            c.alignment = WRAP
            c.border = BORDER
    last = len(rows) + 1
    ws.freeze_panes = freeze
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(last, 2)}"
    names = [h for h, _ in headers]
    if status_col:
        col = get_column_letter(names.index(status_col) + 1)
        dv = DataValidation(type="list", formula1=f'"{",".join(STATUS)}"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"{col}2:{col}{last + 200}")
        rng = f"{col}2:{col}{last + 200}"
        ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Done"'], fill=PatternFill("solid", fgColor="E0ECDF")))
        ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Blocked"'], fill=PatternFill("solid", fgColor="F3DFD9")))
        ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"In progress"'], fill=PatternFill("solid", fgColor="F5EAD2")))
    if priority_col:
        col = get_column_letter(names.index(priority_col) + 1)
        dv = DataValidation(type="list", formula1=f'"{",".join(PRIORITY)}"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"{col}2:{col}{last + 200}")
        ws.conditional_formatting.add(f"{col}2:{col}{last + 200}", CellIsRule(operator="equal", formula=['"Critical"'], font=Font(bold=True, color=ACCENT)))
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    return ws


def summary_sheet(wb: Workbook, source: str, area_col: str, status_col: str, areas: list[str]):
    ws = wb.create_sheet("Summary", 1)
    ws.sheet_view.showGridLines = False
    heads = ["Area", "Items", "Done", "In progress", "Blocked", "Not started", "N/A", "% complete"]
    widths = [38, 9, 9, 12, 10, 12, 8, 12]
    for i, h in enumerate(heads, start=1):
        c = ws.cell(row=1, column=i, value=h)
        c.fill = HEAD_FILL
        c.font = HEAD_FONT
        c.border = BORDER
        ws.column_dimensions[get_column_letter(i)].width = widths[i - 1]
    src = f"'{source}'!"
    for r, a in enumerate(areas, start=2):
        ws.cell(row=r, column=1, value=a)
        ws.cell(row=r, column=2, value=f'=COUNTIF({src}${area_col}:${area_col},A{r})')
        for ci, st in zip(range(3, 8), ["Done", "In progress", "Blocked", "Not started", "N/A"]):
            ws.cell(row=r, column=ci, value=f'=COUNTIFS({src}${area_col}:${area_col},$A{r},{src}${status_col}:${status_col},"{st}")')
        ws.cell(row=r, column=8, value=f"=IFERROR(C{r}/(B{r}-G{r}),0)")
        ws.cell(row=r, column=8).number_format = "0%"
        for ci in range(1, 9):
            ws.cell(row=r, column=ci).border = BORDER
            ws.cell(row=r, column=ci).font = BODY_FONT
    t = len(areas) + 2
    ws.cell(row=t, column=1, value="Total").font = Font(bold=True)
    for ci in range(2, 8):
        col = get_column_letter(ci)
        ws.cell(row=t, column=ci, value=f"=SUM({col}2:{col}{t - 1})").font = Font(bold=True)
    ws.cell(row=t, column=8, value=f"=IFERROR(C{t}/(B{t}-G{t}),0)").number_format = "0%"
    ws.conditional_formatting.add(f"H2:H{t}", CellIsRule(operator="greaterThanOrEqual", formula=["1"], fill=PatternFill("solid", fgColor="E0ECDF")))
    ws.conditional_formatting.add(f"H2:H{t}", CellIsRule(operator="lessThan", formula=["0.5"], fill=PatternFill("solid", fgColor="F3DFD9")))
    ws.cell(row=t + 2, column=1, value="Go-live readiness rule: all Critical items Done or N/A, no Blocked items, ≥ 90% overall.").font = Font(italic=True, size=9, color="676C6E")
    crit = t + 3
    ws.cell(row=crit, column=1, value="Critical items not yet Done")
    ws.cell(row=crit, column=2, value=f'=COUNTIFS({src}$G:$G,"Critical",{src}${status_col}:${status_col},"<>Done",{src}${status_col}:${status_col},"<>N/A",{src}$A:$A,"<>ID")')
    return ws


def save(wb: Workbook, path: str):
    wb.properties.creator = "CloudMigrate.in"
    wb.properties.title = os.path.basename(path)
    wb.save(path)


# ---------------------------------------------------------------- checklists

def pre_items(p: dict) -> list[tuple[str, str, str, str, str]]:
    """(Area, Item, Guidance, Tool / reference, Priority)"""
    return [
        ("Governance & commercials", "Executive sponsor and steering committee named", "Sponsor owns the business case; steering meets fortnightly through migration.", "Programme charter", "Critical"),
        ("Governance & commercials", "Business case approved with run-rate, migration cost and funding shown separately", "Avoid double counting vendor credits against run-rate savings.", p["cost"], "Critical"),
        ("Governance & commercials", "Vendor funding / programme approvals in place", "Confirm eligibility, customer consent and evidence requirements before work starts.", "Funding plan", "High"),
        ("Governance & commercials", "Billing account / subscription ownership agreed", "Who holds the payer / billing root, who pays, and how chargeback works.", p["cost"], "High"),
        ("Governance & commercials", "Commitment strategy agreed (post right-sizing)", "Stage commitments after 30–60 days of post-migration telemetry.", p["commitments"], "Medium"),
        ("Governance & commercials", "Change freeze and blackout calendar captured", "Month-end, quarter-end, festive peaks, audits.", "Wave plan", "High"),
        ("Discovery & assessment", "All in-scope servers discovered with ≥ 14 days of performance data", "Size to p95, not allocation.", p["discovery"], "Critical"),
        ("Discovery & assessment", "Application owners identified and interviewed", "Every application needs a named owner who signs the disposition.", "Application portfolio", "Critical"),
        ("Discovery & assessment", "Dependencies mapped and validated with owners", "Include AD, DNS, NTP, SMTP, file shares, licensing servers, batch schedulers.", p["discovery"], "Critical"),
        ("Discovery & assessment", "7R disposition signed off per application", "Retire, retain, rehost, relocate, replatform, repurchase, refactor.", "Assessment workbook", "High"),
        ("Discovery & assessment", "Database inventory with versions, size, features and growth", "Drives migration method and downtime estimate.", p["db_migration"], "High"),
        ("Discovery & assessment", "Licensing position confirmed (BYOL / licence-included)", "OS, database, middleware and ISV licences — including mobility rights.", "Licence register", "High"),
        ("Discovery & assessment", "Data classification completed for in-scope data", "Personal data under DPDP, regulated data (RBI / SEBI / IRDAI), confidential IP.", "Data classification register", "Critical"),
        ("Landing zone & identity", "Landing zone deployed as code and peer reviewed", p["landing_zone"], p["iac"], "Critical"),
        ("Landing zone & identity", "Guardrails enforced (allowed regions, encryption, tagging, public access)", f"Regions: {p['regions']}.", p["landing_zone"], "Critical"),
        ("Landing zone & identity", "Identity federation and MFA for all administrators", p["identity"], p["identity"], "Critical"),
        ("Landing zone & identity", "Privileged access just-in-time and break-glass accounts tested", "Two break-glass accounts, stored offline, alerting on use.", p["ops_access"], "High"),
        ("Landing zone & identity", "Tagging / labelling standard published and enforced", p["tagging"], p["landing_zone"], "High"),
        ("Network & connectivity", "Private connectivity provisioned and tested for throughput", "Size for replication peaks plus production traffic.", p["connectivity"], "Critical"),
        ("Network & connectivity", "IP address plan with no overlap against on-premises", "Reserve ranges for growth (containers, PaaS private endpoints).", "IPAM", "Critical"),
        ("Network & connectivity", "Hybrid DNS resolution validated both directions", p["dns"], p["dns"], "Critical"),
        ("Network & connectivity", "Firewall rules for migrated apps prepared and reviewed", "Least privilege; document every rule with an owner.", p["network_hub"], "High"),
        ("Network & connectivity", "Internet egress and inspection path designed", "Centralised egress through the hub firewall.", p["network_hub"], "High"),
        ("Network & connectivity", "Load balancer / DNS TTLs lowered ahead of cutover", "Lower TTLs to ≤ 300 s at least 48 h before cutover.", p["dns"], "Medium"),
        ("Security & compliance", "Security baseline enabled in every account / subscription / project", p["posture"], p["benchmark"], "Critical"),
        ("Security & compliance", "Central logging with ≥ 180-day retention in India", "CERT-In directions (April 2022).", p["logging"], "Critical"),
        ("Security & compliance", "Encryption at rest with customer-managed keys where required", p["kms"], p["kms"], "High"),
        ("Security & compliance", "Vulnerability scanning on target images and hosts", p["threat"], p["threat"], "High"),
        ("Security & compliance", "Incident-response runbook updated for cloud (6-hour CERT-In reporting)", "Include contacts, evidence capture and regulator notifications.", "IR runbook", "High"),
        ("Security & compliance", "DPDP obligations reviewed for in-scope personal data", "Notice, consent, purpose limitation, safeguards, breach notification.", "DPDP checklist", "High"),
        ("Data & databases", "Migration method chosen per database with downtime budget", p["db_migration"], p["db_migration"], "Critical"),
        ("Data & databases", "Schema / code conversion assessed for heterogeneous moves", "Stored procedures, triggers, data types, collation and character set.", p["db_migration"], "High"),
        ("Data & databases", "Data validation approach agreed (row counts, checksums, business reports)", "Define pass criteria before migration.", "Validation scripts", "Critical"),
        ("Data & databases", "Bulk data transfer method sized (online vs offline)", p["bulk_transfer"], p["bulk_transfer"], "Medium"),
        ("Application readiness", "Hard-coded IPs, hostnames and file paths identified", "Search configs, scripts, ODBC / JDBC strings and scheduled tasks.", "Code / config scan", "High"),
        ("Application readiness", "Certificates and secrets inventoried and migrated to a vault", p["kms"], p["kms"], "High"),
        ("Application readiness", "Third-party integrations and IP allow-lists notified", "Banks, payment gateways, SaaS, partners that allow-list your egress IPs.", "Integration register", "Critical"),
        ("Application readiness", "OS and runtime support validated on target", "End-of-support OS flagged for remediation or ESU.", "Compatibility matrix", "Medium"),
        ("Application readiness", "Batch jobs and schedulers mapped to new hosts", "Cron, Control-M, SQL Agent, Windows Task Scheduler.", "Batch inventory", "High"),
        ("Backup & DR", "Backup policies defined per tier and applied by default", p["backup"], p["backup"], "Critical"),
        ("Backup & DR", "RTO / RPO agreed per application and DR design approved", p["dr"], p["dr"], "High"),
        ("Backup & DR", "Pre-migration backup of every source system verified restorable", "Your rollback depends on it.", "Backup report", "Critical"),
        ("Testing", "Test plan and UAT scripts signed off by application owners", "Functional, integration, performance and DR tests.", "Test plan", "Critical"),
        ("Testing", "Performance baseline captured on source", "Key transactions, batch runtimes, report durations.", p["monitoring"], "High"),
        ("Testing", "Test migration / test failover completed for the wave", p["replication"], p["replication"], "Critical"),
        ("Cutover readiness", "Cutover runbook rehearsed end-to-end", "Timed rehearsal; update durations.", "Cutover runbook", "Critical"),
        ("Cutover readiness", "Rollback triggers and decision owner agreed", "Who calls rollback, at what time, on what evidence.", "Cutover runbook", "Critical"),
        ("Cutover readiness", "Go / no-go meeting scheduled with quorum", "Business, IT, security, vendor.", "Go / no-go checklist", "High"),
        ("Cutover readiness", "Replication healthy and lag within threshold", p["replication"], p["replication"], "Critical"),
        ("Communications", "User and stakeholder communications sent", "Outage window, what changes, who to call.", "Comms plan", "High"),
        ("Communications", "Service desk briefed with known issues and FAQs", "Temporary hypercare queue created.", "ITSM", "Medium"),
        ("Operations readiness", "Monitoring and alerting configured on target before cutover", p["monitoring"], p["monitoring"], "Critical"),
        ("Operations readiness", "Patching schedule and maintenance windows defined", p["patching"], p["patching"], "High"),
        ("Operations readiness", "Runbooks and RACI handed to the operations team", "Run-team shadows the final rehearsal.", "Runbook & RACI", "High"),
        ("Operations readiness", "Cost budgets and anomaly alerts configured", p["cost"], p["cost"], "Medium"),
    ]


def post_items(p: dict) -> list[tuple[str, str, str, str, str]]:
    return [
        ("Functional validation", "Application smoke tests passed", "Log-in, key transactions, reports, integrations.", "UAT scripts", "Critical"),
        ("Functional validation", "End-to-end business process tested by the owner", "Owner signs the result, not the project team.", "UAT sign-off", "Critical"),
        ("Functional validation", "Integrations and interfaces confirmed (inbound and outbound)", "Payment gateways, SaaS, partners, file drops.", "Integration register", "Critical"),
        ("Functional validation", "Batch jobs ran successfully on schedule", "Compare runtime with the source baseline.", "Scheduler logs", "High"),
        ("Functional validation", "Email / SMTP relay, printing and file shares working", "Often forgotten until Monday morning.", "Smoke tests", "Medium"),
        ("Performance", "Key transaction response times within baseline tolerance", "Typical tolerance ± 10%.", p["monitoring"], "High"),
        ("Performance", "Resource utilisation reviewed after 7 days", "Right-size over- and under-provisioned resources.", p["monitoring"], "Medium"),
        ("Performance", "Database wait events / query performance reviewed", "Compare top SQL with the source.", "DB monitoring", "High"),
        ("Data integrity", "Row counts and checksums reconciled", "All tables or a signed sample for very large tables.", "Validation scripts", "Critical"),
        ("Data integrity", "Business control totals match (e.g. ledger balances)", "Signed by finance or the data owner.", "Reconciliation report", "Critical"),
        ("Data integrity", "Replication / CDC stopped and source set read-only", "Prevents split-brain writes.", p["db_migration"], "Critical"),
        ("Security", "Security posture score reviewed for new resources", p["posture"], p["benchmark"], "High"),
        ("Security", "No public exposure beyond design (ports, buckets, endpoints)", "Scan external attack surface.", p["posture"], "Critical"),
        ("Security", "Logs flowing to central logging and SIEM", p["siem"], p["logging"], "Critical"),
        ("Security", "Privileged access reviewed; temporary migration access removed", "Remove migration service accounts and keys.", p["identity"], "Critical"),
        ("Security", "Encryption and key ownership verified", p["kms"], p["kms"], "High"),
        ("Monitoring", "Alerts tested end-to-end (fire, route, acknowledge)", p["monitoring"], p["monitoring"], "High"),
        ("Monitoring", "Dashboards published for the run team and owners", p["monitoring"], p["monitoring"], "Medium"),
        ("Backup & DR", "First backups completed and a restore tested", p["backup"], p["backup"], "Critical"),
        ("Backup & DR", "DR replication configured and failover test scheduled", p["dr"], p["dr"], "High"),
        ("Cost & tagging", "All resources tagged per standard", p["tagging"], p["landing_zone"], "High"),
        ("Cost & tagging", "Cost allocated to the right cost centre in reports", p["cost"], p["cost"], "Medium"),
        ("Cost & tagging", "Commitment / reservation candidates identified after 30 days", p["commitments"], p["cost"], "Medium"),
        ("Cost & tagging", "Vendor funding evidence captured (tags, IDs, reports)", "Subscription / account IDs, migration reports, customer attestation.", "Funding tracker", "High"),
        ("Operations handover", "Runbooks, RACI and SLAs accepted by the run team", "Signed handover form.", "Runbook & RACI", "Critical"),
        ("Operations handover", "Patching and maintenance windows active", p["patching"], p["patching"], "High"),
        ("Operations handover", "CMDB / asset register updated", "Retire source CIs, add target CIs.", "CMDB", "Medium"),
        ("Operations handover", "Hypercare exit criteria met", "No open P1 / P2, stable for 10 business days.", "Hypercare log", "High"),
        ("Decommission", "Source systems powered off after agreed retention period", "Typically 2–4 weeks after successful cutover.", "Decommission plan", "High"),
        ("Decommission", "Source data securely wiped / media destroyed with certificate", "Required for regulated data.", "Destruction certificate", "High"),
        ("Decommission", "Licences, support contracts and hosting released", "Capture savings in the benefits tracker.", "Contract register", "Medium"),
        ("Documentation", "As-built architecture and network diagrams updated", "Stored in the design repository.", "Architecture repo", "Medium"),
        ("Documentation", "Lessons learned captured for the next wave", "Update runbook durations and risks.", "Retrospective", "Medium"),
        ("Business sign-off", "Application owner acceptance signed", "Formal closure for the wave.", "Sign-off form", "Critical"),
        ("Business sign-off", "Benefits tracker updated (cost, performance, risk)", "Feeds the business case review.", "Benefits tracker", "Medium"),
    ]


def build_checklist(p: dict, slug: str, kind: str):
    items = pre_items(p) if kind == "pre" else post_items(p)
    title = "Pre-Migration Checklist" if kind == "pre" else "Post-Migration Checklist"
    wb = Workbook()
    cover_sheet(
        wb,
        f"{title} — {p['short']}",
        p,
        "Readiness gate before each migration wave." if kind == "pre" else "Validation and handover gate after each migration wave.",
        [
            "Copy the Checklist sheet once per wave (or filter by wave in the Wave column).",
            "Assign an owner and due date to every item; set Status from the drop-down.",
            "Attach evidence links in the Evidence column — auditors and funding reviewers will ask.",
            "The Summary sheet calculates completion by area; go-live needs every Critical item Done or N/A.",
        ],
    )
    rows = []
    for i, (area, item, guide, tool, prio) in enumerate(items, start=1):
        rows.append([f"{'PRE' if kind == 'pre' else 'POST'}-{i:03d}", area, item, guide, tool, "", prio, "Not started", "", "", ""])
    table_sheet(
        wb,
        "Checklist",
        [("ID", 10), ("Area", 24), ("Item", 46), ("Guidance", 52), (f"{p['short']} tool / reference", 40), ("Owner", 18), ("Priority", 11), ("Status", 13), ("Due date", 12), ("Wave", 10), ("Evidence / notes", 40)],
        rows,
        status_col="Status",
        priority_col="Priority",
        freeze="D2",
    )
    areas = list(dict.fromkeys(a for a, *_ in items))
    summary_sheet(wb, "Checklist", "B", "H", areas)
    fname = f"{slug}-{'pre' if kind == 'pre' else 'post'}-migration-checklist.xlsx"
    save(wb, os.path.join(OUT, slug, fname))
    return fname, len(items)


# ---------------------------------------------------------------- assessment workbook

def build_assessment(p: dict, slug: str):
    wb = Workbook()
    cover_sheet(
        wb,
        f"Discovery & Assessment Workbook — {p['short']}",
        p,
        "Capture inventory, decide dispositions, plan waves and build the TCO.",
        [
            f"Populate Server Inventory from {p['discovery']} exports.",
            "Complete the Application Portfolio with owners; score each app in the 7R Matrix.",
            "Record dependencies to form move groups; assign waves in the Wave Plan.",
            "Enter on-premises and target costs in TCO — totals and savings calculate automatically.",
        ],
    )
    inv = table_sheet(
        wb,
        "Server Inventory",
        [("Server", 18), ("Application", 22), ("Environment", 13), ("OS / version", 22), ("vCPU", 7), ("RAM (GB)", 9), ("Storage (GB)", 11), ("CPU p95 %", 10), ("RAM p95 %", 10), ("Target size", 18), ("Disposition", 13), ("Wave", 8), ("Notes", 30)],
        [
            ["APP-WEB-01", "Customer portal", "Production", "Windows Server 2016", 4, 16, 200, 38, 61, "", "Rehost", "W2", "Behind F5"],
            ["APP-DB-01", "Customer portal", "Production", "Linux / PostgreSQL 13", 8, 64, 1200, 55, 72, "", "Replatform", "W2", "Managed PostgreSQL"],
            ["FIN-ERP-01", "Finance ERP", "Production", "Oracle Linux 7 / Oracle 19c", 16, 128, 4000, 47, 80, "", "Rehost", "W4", "Licence BYOL"],
            ["HR-APP-01", "HRMS (legacy)", "Production", "Windows Server 2012 R2", 2, 8, 120, 6, 25, "", "Repurchase", "—", "Replace with SaaS"],
        ],
        freeze="B2",
    )
    dv = DataValidation(type="list", formula1='"Retire,Retain,Rehost,Relocate,Replatform,Repurchase,Refactor"', allow_blank=True)
    inv.add_data_validation(dv)
    dv.add("K2:K2000")
    port = table_sheet(
        wb,
        "Application Portfolio",
        [("Application", 24), ("Owner", 18), ("Business criticality (1-5)", 12), ("Data classification", 16), ("Users", 9), ("Change frequency (1-5)", 12), ("Technical debt (1-5)", 12), ("Disposition", 13), (f"Target {p['short']} service", 30), ("Complexity score", 11), ("Wave", 8), ("Notes", 30)],
        [
            ["Customer portal", "Head of Digital", 5, "Personal (DPDP)", 120000, 4, 3, "Replatform", "", "", "W2", ""],
            ["Finance ERP", "CFO office", 5, "Confidential", 300, 2, 2, "Rehost", "", "", "W4", "Month-end freeze"],
            ["HRMS (legacy)", "CHRO office", 3, "Personal (DPDP)", 2500, 1, 5, "Repurchase", "SaaS", "", "—", "Contract ends Mar"],
        ],
        freeze="B2",
    )
    for r in range(2, 202):
        port.cell(row=r, column=10, value=f'=IF(A{r}="","",ROUND(C{r}*0.4+F{r}*0.2+G{r}*0.4,1))')
    port.add_data_validation(dv2 := DataValidation(type="list", formula1='"Retire,Retain,Rehost,Relocate,Replatform,Repurchase,Refactor"', allow_blank=True))
    dv2.add("H2:H2000")
    m = table_sheet(
        wb,
        "7R Decision Matrix",
        [("Criterion", 34), ("Weight", 8), ("Rehost", 9), ("Relocate", 9), ("Replatform", 10), ("Refactor", 9), ("Repurchase", 10), ("Retire / Retain", 12)],
        [
            ["Speed to exit data centre", 0.25, 5, 5, 3, 1, 3, 5],
            ["Run-cost reduction", 0.20, 2, 2, 4, 5, 4, 3],
            ["Risk / complexity (5 = low risk)", 0.20, 5, 4, 3, 2, 3, 5],
            ["Business agility gained", 0.15, 1, 1, 3, 5, 4, 1],
            ["Skills available", 0.10, 5, 4, 3, 2, 4, 5],
            ["Licence impact", 0.10, 3, 3, 4, 4, 3, 5],
        ],
        freeze="B2",
    )
    m.cell(row=8, column=1, value="Weighted score").font = Font(bold=True)
    for ci in range(3, 9):
        col = get_column_letter(ci)
        m.cell(row=8, column=ci, value=f"=SUMPRODUCT($B$2:$B$7,{col}2:{col}7)").font = Font(bold=True)
    m.cell(row=10, column=1, value="Adjust weights per application class; weights should total 1.0.").font = Font(italic=True, size=9)
    m.cell(row=11, column=1, value="Weight total").font = Font(size=9)
    m.cell(row=11, column=2, value="=SUM(B2:B7)")
    table_sheet(
        wb,
        "Dependencies",
        [("Source application", 24), ("Source server", 16), ("Target application", 24), ("Target server", 16), ("Port", 8), ("Protocol", 10), ("Direction", 11), ("Criticality", 11), ("Move group", 12), ("Validated by owner?", 14)],
        [["Customer portal", "APP-WEB-01", "Customer portal", "APP-DB-01", 5432, "TCP", "Outbound", "High", "MG-01", "Yes"], ["Finance ERP", "FIN-ERP-01", "Active Directory", "DC-01", 389, "TCP", "Outbound", "High", "Shared", "Yes"]],
    )
    tco = table_sheet(
        wb,
        "TCO (3-year)",
        [("Cost line (INR)", 40), ("Year 1", 14), ("Year 2", 14), ("Year 3", 14), ("Total", 14), ("Notes", 40)],
        [
            ["ON-PREMISES", None, None, None, None, ""],
            ["Hardware refresh / depreciation", 0, 0, 0, None, ""],
            ["Data-centre space, power, cooling", 0, 0, 0, None, ""],
            ["Software licences & support", 0, 0, 0, None, ""],
            ["Operations staff (allocated)", 0, 0, 0, None, ""],
            ["Network / connectivity", 0, 0, 0, None, ""],
            ["DR site", 0, 0, 0, None, ""],
            ["On-premises total", None, None, None, None, ""],
            [f"{p['short'].upper()} TARGET", None, None, None, None, ""],
            ["Compute (after right-sizing)", 0, 0, 0, None, p["commitments"]],
            ["Storage & backup", 0, 0, 0, None, ""],
            ["Databases / PaaS", 0, 0, 0, None, ""],
            ["Network, egress, connectivity", 0, 0, 0, None, p["connectivity"]],
            ["Security & monitoring services", 0, 0, 0, None, ""],
            ["Licences (BYOL / included)", 0, 0, 0, None, ""],
            ["Managed services", 0, 0, 0, None, ""],
            ["Cloud total (run-rate)", None, None, None, None, ""],
            ["ONE-TIME", None, None, None, None, ""],
            ["Migration services", 0, 0, 0, None, ""],
            ["Dual-running during migration", 0, 0, 0, None, ""],
            ["Vendor funding / credits (enter as negative)", 0, 0, 0, None, "Keep separate from run-rate"],
            ["One-time total", None, None, None, None, ""],
            ["NET SAVING (on-prem − cloud − one-time)", None, None, None, None, ""],
        ],
        freeze="B2",
    )
    # Sheet rows: section headers 2 / 10 / 19; cost lines 3–8, 11–17, 20–22; totals 9 / 18 / 23; net 24.
    for r in [*range(3, 9), *range(11, 18), *range(20, 23)]:
        tco.cell(row=r, column=5, value=f"=SUM(B{r}:D{r})")
    for ci in range(2, 6):
        col = get_column_letter(ci)
        tco.cell(row=9, column=ci, value=f"=SUM({col}3:{col}8)").font = Font(bold=True)
        tco.cell(row=18, column=ci, value=f"=SUM({col}11:{col}17)").font = Font(bold=True)
        tco.cell(row=23, column=ci, value=f"=SUM({col}20:{col}22)").font = Font(bold=True)
        tco.cell(row=24, column=ci, value=f"={col}9-{col}18-{col}23").font = Font(bold=True, color=ACCENT)
    for r in range(3, 25):
        for ci in range(2, 6):
            tco.cell(row=r, column=ci).number_format = "#,##0;[Red]-#,##0"
    for r in (9, 18, 23, 24):
        tco.cell(row=r, column=1).font = Font(bold=True, color=INK)
    for r in (2, 10, 19):
        tco.cell(row=r, column=1).font = Font(bold=True, color=ACCENT)
        for ci in range(2, 6):
            tco.cell(row=r, column=ci).value = None
    table_sheet(
        wb,
        "Wave Plan",
        [("Wave", 8), ("Applications", 40), ("Servers", 9), ("Databases", 10), ("Start", 12), ("Cutover date", 13), ("Downtime window", 16), ("Rollback deadline", 16), ("Lead", 16), ("Status", 13)],
        [["W1 (pilot)", "Intranet, file share", 4, 0, "", "", "Sat 22:00–02:00", "Sun 01:00", "", "Not started"], ["W2", "Customer portal", 6, 1, "", "", "Sat 22:00–04:00", "Sun 02:00", "", "Not started"]],
        status_col="Status",
    )
    fname = f"{slug}-discovery-assessment-workbook.xlsx"
    save(wb, os.path.join(OUT, slug, fname))
    return fname


# ---------------------------------------------------------------- security checklist

def security_items(p: dict):
    b = p["benchmark_short"]
    return [
        ("Identity & access", "MFA enforced for all human administrators", p["identity"], f"{b} IAM", "DPDP s.8(5) safeguards", "Critical"),
        ("Identity & access", "No standing privileged access; just-in-time elevation", p["ops_access"], f"{b} IAM", "RBI cyber-resilience", "Critical"),
        ("Identity & access", "Root / global admin / super-admin accounts locked down and alerted", "Hardware MFA, no access keys, alert on use.", f"{b} IAM", "", "Critical"),
        ("Identity & access", "Break-glass accounts documented and tested quarterly", "", f"{b} IAM", "", "High"),
        ("Identity & access", "Service identities use workload identity, not long-lived keys", "Rotate or remove keys older than 90 days.", f"{b} IAM", "", "High"),
        ("Identity & access", "Access reviews performed quarterly for privileged roles", "", "ISO 27001 A.5.18", "SEBI cloud framework", "High"),
        ("Identity & access", "Guest / external identities reviewed and expired", "", f"{b} IAM", "", "Medium"),
        ("Logging & monitoring", "Control-plane audit logs enabled organisation-wide", p["logging"], f"{b} Logging", "CERT-In 2022", "Critical"),
        ("Logging & monitoring", "Logs retained ≥ 180 days in India and protected from deletion", p["logging"], f"{b} Logging", "CERT-In 2022", "Critical"),
        ("Logging & monitoring", "Time synchronisation to approved NTP sources", "CERT-In expects NIC / NPL NTP sync for ICT systems.", "CIS v8 8.4", "CERT-In 2022", "High"),
        ("Logging & monitoring", "Threat detection enabled in all regions and accounts", p["threat"], f"{b} Monitoring", "", "Critical"),
        ("Logging & monitoring", "Security posture management enabled with a tracked score", p["posture"], b, "", "High"),
        ("Logging & monitoring", "SIEM ingests identity, control-plane, network and endpoint logs", p["siem"], "CIS v8 8.9", "", "High"),
        ("Logging & monitoring", "Alerts for root / admin use, policy changes, key deletion", "", f"{b} Monitoring", "", "High"),
        ("Data protection", "Encryption at rest everywhere; customer-managed keys for regulated data", p["kms"], f"{b} Storage", "RBI / IRDAI", "Critical"),
        ("Data protection", "Encryption in transit (TLS 1.2+) enforced", "", "CIS v8 3.10", "", "High"),
        ("Data protection", "No public object storage unless explicitly approved", "", f"{b} Storage", "DPDP", "Critical"),
        ("Data protection", "Personal-data discovery and classification running", p["data_protection"], "CIS v8 3.2", "DPDP", "High"),
        ("Data protection", "Key rotation and separation of key administrators", p["kms"], f"{b} KMS", "", "Medium"),
        ("Data protection", "Non-production data masked or synthetic", "", "ISO 27001 A.8.11", "DPDP", "High"),
        ("Network security", "No management ports (22 / 3389) open to the internet", "Use bastion / session manager.", f"{b} Networking", "", "Critical"),
        ("Network security", "Default-deny network policies and segmentation between tiers", p["network_hub"], f"{b} Networking", "", "High"),
        ("Network security", "WAF and DDoS protection on internet-facing applications", p["waf"], "CIS v8 13.10", "", "High"),
        ("Network security", "Private endpoints for PaaS data services", "", f"{b} Networking", "", "High"),
        ("Network security", "Egress filtered through central inspection", p["network_hub"], "CIS v8 13", "", "Medium"),
        ("Network security", "Flow logs enabled on critical networks", "", f"{b} Logging", "CERT-In 2022", "Medium"),
        ("Workload security", "Vulnerability scanning of hosts, containers and images", p["threat"], "CIS v8 7", "", "High"),
        ("Workload security", "Critical patches applied within 14 days", p["patching"], "CIS v8 7.3", "RBI", "High"),
        ("Workload security", "Endpoint detection and response on all servers", p["threat"], "CIS v8 10", "", "High"),
        ("Workload security", "Hardened images (CIS level 1) used as baselines", "", "CIS v8 4.1", "", "Medium"),
        ("Workload security", "Secrets stored in a vault, not code or config", p["kms"], "CIS v8 3.11", "", "High"),
        ("Resilience", "Immutable / locked backups for tier-1 data", p["backup"], "CIS v8 11", "RBI", "Critical"),
        ("Resilience", "Restore tests performed and documented", "", "CIS v8 11.5", "SEBI", "High"),
        ("Resilience", "DR plan tested within the last 12 months", p["dr"], "ISO 27001 A.5.30", "RBI / SEBI", "High"),
        ("Governance", "Guardrails / policies prevent non-compliant deployments", p["landing_zone"], b, "", "High"),
        ("Governance", "Resources restricted to approved Indian regions where residency applies", p["regions"], b, "DPDP / sector rules", "High"),
        ("Governance", "Asset inventory complete and owners assigned", "", "CIS v8 1", "", "Medium"),
        ("Governance", "Third-party and support access logged and time-bound", "", "ISO 27001 A.5.19", "RBI outsourcing", "High"),
        ("Incident response", "Cloud IR playbooks for credential compromise, data exposure, ransomware", "", "CIS v8 17", "CERT-In 2022", "High"),
        ("Incident response", "Ability to report to CERT-In within 6 hours demonstrated", "Contacts, templates, evidence capture.", "", "CERT-In 2022", "Critical"),
        ("Incident response", "DPDP breach notification workflow to the Board and data principals", "", "", "DPDP Act / Rules", "High"),
    ]


def build_security(p: dict, slug: str):
    wb = Workbook()
    cover_sheet(
        wb,
        f"Cloud Security Assessment Checklist — {p['short']}",
        p,
        f"Assess {p['short']} posture against {p['benchmark']} and Indian regulatory obligations.",
        [
            "Work area by area; record Result (Pass / Partial / Fail) and evidence for each control.",
            "Severity reflects impact if the control fails; prioritise Critical and High fails.",
            "The Regulation column shows the Indian obligation most affected — validate with your compliance team.",
            "The Summary sheet tracks progress; export failing rows to your remediation tracker.",
        ],
    )
    rows = []
    for i, (area, ctl, svc, ref, reg, sev) in enumerate(security_items(p), start=1):
        rows.append([f"SEC-{i:03d}", area, ctl, svc, ref, reg, sev, "Not started", "", "", "", ""])
    ws = table_sheet(
        wb,
        "Controls",
        [("ID", 9), ("Area", 18), ("Control", 48), (f"{p['short']} service / method", 40), ("Benchmark ref", 16), ("Regulation", 18), ("Priority", 11), ("Status", 13), ("Result", 11), ("Evidence", 30), ("Owner", 14), ("Remediation / due", 28)],
        rows,
        status_col="Status",
        priority_col="Priority",
        freeze="D2",
    )
    dv = DataValidation(type="list", formula1='"Pass,Partial,Fail,Not assessed"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add("I2:I500")
    ws.conditional_formatting.add("I2:I500", CellIsRule(operator="equal", formula=['"Fail"'], fill=PatternFill("solid", fgColor="F3DFD9")))
    ws.conditional_formatting.add("I2:I500", CellIsRule(operator="equal", formula=['"Pass"'], fill=PatternFill("solid", fgColor="E0ECDF")))
    areas = list(dict.fromkeys(r[1] for r in rows))
    summary_sheet(wb, "Controls", "B", "H", areas)
    fname = f"{slug}-security-assessment-checklist.xlsx"
    save(wb, os.path.join(OUT, slug, fname))
    return fname, len(rows)


# ---------------------------------------------------------------- managed services runbook & RACI

def build_runbook(p: dict, slug: str):
    wb = Workbook()
    cover_sheet(
        wb,
        f"Managed Services Runbook, RACI & SLA — {p['short']}",
        p,
        "Define who does what, how fast, and how the service is measured.",
        ["Tailor roles to your organisation in the RACI sheet.", "Agree SLA targets and escalation contacts before transition.", "Use the Operational Runbook as the schedule of recurring tasks.", "Review KPIs monthly in the service review."],
    )
    acts = [
        "Incident detection & triage", "Major incident management", "Problem management / RCA", "Standard change execution", "Normal change approval (CAB)",
        "OS patching", "Database patching", "Backup monitoring", "Restore requests & tests", "DR drills", "Security alert triage", "Vulnerability remediation",
        "Access requests & reviews", "Capacity & performance management", "Cost optimisation / FinOps", "Commitment / reservation management", "Landing-zone policy changes",
        "Network & firewall changes", "Certificate renewals", "Monthly service review", "Regulatory evidence & audits",
    ]
    roles = ["Customer IT lead", "Application owner", "CloudMigrate service manager", "CloudMigrate L1/L2 ops", "CloudMigrate L3 architect", "Security lead", "Vendor support"]
    matrix = {
        "Incident detection & triage": "I I A R C I C", "Major incident management": "A C R R R C C", "Problem management / RCA": "I C A R R C C",
        "Standard change execution": "I I A R C I I", "Normal change approval (CAB)": "A C R C C C I", "OS patching": "I C A R C I I",
        "Database patching": "I C A R R I C", "Backup monitoring": "I I A R I I I", "Restore requests & tests": "C R A R C I I",
        "DR drills": "A R R R R C C", "Security alert triage": "I I A R C R I", "Vulnerability remediation": "C C A R R R I",
        "Access requests & reviews": "A R R R I C I", "Capacity & performance management": "I C A R R I I", "Cost optimisation / FinOps": "A C R R R I I",
        "Commitment / reservation management": "A I R C R I I", "Landing-zone policy changes": "A I R C R C I", "Network & firewall changes": "A C R R R C I",
        "Certificate renewals": "I C A R I I I", "Monthly service review": "A C R C C C I", "Regulatory evidence & audits": "A C R R C R I",
    }
    rows = [[a] + matrix[a].split() for a in acts]
    ws = table_sheet(wb, "RACI", [("Activity", 36)] + [(r, 16) for r in roles], rows, freeze="B2")
    dv = DataValidation(type="list", formula1='"R,A,C,I,R/A"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"B2:{get_column_letter(len(roles) + 1)}200")
    for ci in range(2, len(roles) + 2):
        for r in range(2, len(rows) + 2):
            ws.cell(row=r, column=ci).alignment = Alignment(horizontal="center", vertical="center")
    ws.conditional_formatting.add(f"B2:{get_column_letter(len(roles) + 1)}200", CellIsRule(operator="equal", formula=['"A"'], fill=PatternFill("solid", fgColor="F6E2DA"), font=Font(bold=True, color=ACCENT)))
    ws.conditional_formatting.add(f"B2:{get_column_letter(len(roles) + 1)}200", CellIsRule(operator="equal", formula=['"R"'], fill=PatternFill("solid", fgColor="DDE9E7")))
    ws.cell(row=len(rows) + 3, column=1, value="R = Responsible · A = Accountable (exactly one) · C = Consulted · I = Informed").font = Font(italic=True, size=9)
    table_sheet(
        wb,
        "SLA Matrix",
        [("Priority", 10), ("Definition", 50), ("Response", 12), ("Restore target", 14), ("Update frequency", 16), ("Hours", 12)],
        [
            ["P1", "Production down or security incident with data at risk; no workaround.", "15 min", "4 h", "Every 30 min", "24×7"],
            ["P2", "Production degraded or a critical function unavailable; workaround exists.", "30 min", "8 h", "Every 2 h", "24×7"],
            ["P3", "Non-critical function impaired or single user affected.", "4 h", "3 business days", "Daily", "Business hours"],
            ["P4", "Request, question or cosmetic issue.", "1 business day", "10 business days", "On resolution", "Business hours"],
        ],
    )
    table_sheet(
        wb,
        "Operational Runbook",
        [("Task", 36), ("Frequency", 14), (f"{p['short']} tool", 42), ("Procedure summary", 56), ("Owner", 18), ("Evidence", 22)],
        [
            ["Review overnight alerts and backup jobs", "Daily", f"{p['monitoring']}; {p['backup']}", "Check failed jobs, re-run, raise incidents for repeated failures.", "L1 ops", "Daily ops log"],
            ["Security findings triage", "Daily", f"{p['posture']}; {p['siem']}", "Triage new high / critical findings within 24 h; assign owners.", "Security analyst", "SIEM case notes"],
            ["Patch deployment", "Monthly (+ emergency)", p["patching"], "Pilot ring → production rings; verify compliance ≥ 95% within 14 days.", "L2 ops", "Patch report"],
            ["Restore test", "Monthly", p["backup"], "Restore a tier-1 system to an isolated network; verify application start.", "L2 ops + app owner", "Restore test record"],
            ["Cost review", "Monthly", p["cost"], "Idle resources, right-sizing, commitments coverage, anomalies; actions with owners.", "FinOps lead", "Cost review deck"],
            ["Access review", "Quarterly", p["identity"], "Review privileged roles and remove stale access.", "Security lead", "Signed review"],
            ["DR drill", "Semi-annual", p["dr"], "Planned failover of a tier-1 service; measure RTO / RPO.", "L3 architect", "DR report"],
            ["Architecture / well-architected review", "Quarterly", p["landing_zone"], "Review reliability, security, cost, performance and operations; update backlog.", "L3 architect", "Review report"],
            ["Certificate expiry check", "Weekly", "Certificate inventory", "Renew certificates expiring in < 30 days.", "L2 ops", "Cert register"],
            ["Log retention verification", "Quarterly", p["logging"], "Confirm ≥ 180-day in-country retention and immutability.", "Security lead", "Retention evidence"],
        ],
    )
    table_sheet(
        wb,
        "Escalation Matrix",
        [("Level", 10), ("Role", 28), ("Name", 22), ("Phone", 16), ("Email", 28), ("Escalate after", 18)],
        [["L1", "Service desk / NOC", "", "", "", "—"], ["L2", "Platform engineer on call", "", "", "", "30 min (P1)"], ["L3", "Lead architect", "", "", "", "1 h (P1)"], ["L4", "Service delivery manager", "", "", "", "2 h (P1)"], ["L5", "Account executive / customer CIO", "", "", "", "4 h (P1)"]],
    )
    table_sheet(
        wb,
        "Monthly KPIs",
        [("KPI", 40), ("Target", 16), ("Actual", 12), ("Trend", 10), ("Commentary", 40)],
        [
            ["P1 response within SLA", "100%", "", "", ""], ["P2 restore within SLA", "≥ 95%", "", "", ""], ["Patch compliance (14 days)", "≥ 95%", "", "", ""],
            ["Backup success rate", "≥ 99%", "", "", ""], ["Restore tests passed", "100%", "", "", ""], ["Security posture score", "+ trend", "", "", ""],
            ["Critical findings > 30 days", "0", "", "", ""], ["Commitment coverage", "≥ 70%", "", "", ""], ["Unallocated (untagged) spend", "< 5%", "", "", ""], ["Changes causing incidents", "< 2%", "", "", ""],
        ],
    )
    fname = f"{slug}-managed-services-runbook-raci.xlsx"
    save(wb, os.path.join(OUT, slug, fname))
    return fname


# ---------------------------------------------------------------- cutover runbook

def build_cutover(p: dict, slug: str):
    wb = Workbook()
    cover_sheet(
        wb,
        f"Cutover Runbook — {p['short']}",
        p,
        "Minute-by-minute plan for a migration wave cutover with go / no-go and rollback.",
        ["Set the planned start time in cell C2 of the Cutover Plan sheet; start / end times calculate from durations.", "Rehearse at least once and update durations with real timings.", "Hold the go / no-go meeting with quorum; record the decision.", "If a rollback trigger fires, switch to the Rollback Plan sheet immediately."],
    )
    ws = wb.create_sheet("Cutover Plan")
    ws.sheet_view.showGridLines = False
    ws["A1"] = "Planned start"
    ws["A1"].font = Font(bold=True)
    ws["C2"] = dt.datetime(2026, 11, 7, 22, 0)
    ws["C2"].number_format = "dd-mmm-yyyy hh:mm"
    ws["A2"] = "Cutover start (edit) →"
    heads = [("Step", 6), ("Phase", 16), ("Task", 52), ("Owner", 16), ("Duration (min)", 10), ("Planned start", 17), ("Planned end", 17), ("Go/No-go gate", 10), ("Status", 13), ("Actual end", 17), ("Rollback ref", 10), ("Notes", 30)]
    for i, (h, w) in enumerate(heads, start=1):
        c = ws.cell(row=4, column=i, value=h)
        c.fill = HEAD_FILL
        c.font = HEAD_FONT
        c.border = BORDER
        c.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = w
    steps = [
        ("T-7 days", "Confirm wave scope, owners, communications and change approval", "Migration lead", 0, "", "RB-0"),
        ("T-2 days", "Lower DNS TTLs to ≤ 300 s; confirm replication healthy and lag < threshold", "Network lead", 0, "", "RB-0"),
        ("T-1 day", "Final test migration / failover reviewed; known issues logged", "Migration engineer", 0, "", "RB-0"),
        ("Go / no-go", "Go / no-go meeting: readiness checklist, staffing, rollback owner present", "Migration lead", 30, "GATE", "RB-0"),
        ("Freeze", "Announce outage; disable user access / put application in maintenance mode", "App owner", 15, "", "RB-1"),
        ("Freeze", "Stop application services and scheduled jobs on source", "App support", 15, "", "RB-1"),
        ("Data sync", "Confirm final replication / CDC sync; record source row counts and checksums", "DBA", 30, "", "RB-2"),
        ("Data sync", "Set source databases read-only", "DBA", 10, "", "RB-2"),
        ("Cutover", f"Launch / promote target instances ({p['replication']})", "Migration engineer", 30, "", "RB-3"),
        ("Cutover", f"Complete database switchover ({p['db_migration']})", "DBA", 30, "", "RB-3"),
        ("Cutover", "Update connection strings, secrets and configuration", "App support", 20, "", "RB-3"),
        ("Cutover", "Switch DNS / load balancer to target", "Network lead", 15, "", "RB-4"),
        ("Validate", "Technical smoke tests (services up, ports, logs, monitoring)", "Migration engineer", 20, "", "RB-4"),
        ("Validate", "Data reconciliation: row counts, checksums, control totals", "DBA + data owner", 30, "", "RB-4"),
        ("Validate", "Business validation by key users", "App owner", 45, "GATE", "RB-4"),
        ("Decision", "Go-live decision or rollback (before rollback deadline)", "Migration lead", 15, "GATE", "RB-5"),
        ("Go-live", "Re-enable user access, scheduled jobs and integrations", "App support", 15, "", ""),
        ("Go-live", "Communicate completion; open hypercare queue", "Migration lead", 10, "", ""),
        ("Hypercare", "Monitor for 72 h; daily check-ins with owners", "Ops lead", 0, "", ""),
    ]
    for i, (phase, task, owner, dur, gate, rb) in enumerate(steps, start=5):
        ws.cell(row=i, column=1, value=i - 4)
        ws.cell(row=i, column=2, value=phase)
        ws.cell(row=i, column=3, value=task)
        ws.cell(row=i, column=4, value=owner)
        ws.cell(row=i, column=5, value=dur)
        if i == 5:
            ws.cell(row=i, column=6, value="=$C$2")
        else:
            ws.cell(row=i, column=6, value=f"=G{i - 1}")
        ws.cell(row=i, column=7, value=f"=F{i}+E{i}/1440")
        ws.cell(row=i, column=8, value=gate)
        ws.cell(row=i, column=9, value="Not started")
        ws.cell(row=i, column=11, value=rb)
        for ci in range(1, 13):
            c = ws.cell(row=i, column=ci)
            c.font = BODY_FONT
            c.border = BORDER
            c.alignment = WRAP
        ws.cell(row=i, column=6).number_format = "dd-mmm hh:mm"
        ws.cell(row=i, column=7).number_format = "dd-mmm hh:mm"
        ws.cell(row=i, column=10).number_format = "dd-mmm hh:mm"
    last = 4 + len(steps)
    # The first three rows are T-minus activities and should not consume the window
    for r in (5, 6, 7):
        ws.cell(row=r, column=6, value="")
        ws.cell(row=r, column=7, value="")
    ws.cell(row=8, column=6, value="=$C$2")
    dv = DataValidation(type="list", formula1=f'"{",".join(STATUS)}"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"I5:I{last + 50}")
    ws.conditional_formatting.add(f"A5:L{last}", FormulaRule(formula=['$H5="GATE"'], fill=PatternFill("solid", fgColor="F6E2DA")))
    ws.conditional_formatting.add(f"I5:I{last + 50}", CellIsRule(operator="equal", formula=['"Done"'], fill=PatternFill("solid", fgColor="E0ECDF")))
    ws.freeze_panes = "D5"
    ws.cell(row=last + 2, column=3, value="Total window (hours)").font = Font(bold=True)
    ws.cell(row=last + 2, column=5, value=f"=SUM(E8:E{last})/60").number_format = "0.0"
    table_sheet(
        wb,
        "Go-No-Go",
        [("Criterion", 56), ("Owner", 18), ("Met?", 8), ("Evidence", 36)],
        [
            ["Pre-migration checklist: all Critical items Done or N/A", "Migration lead", "", ""],
            ["Test migration / failover passed in the last 7 days", "Migration engineer", "", ""],
            ["Replication lag within threshold for 24 h", "Migration engineer", "", ""],
            ["Verified backup of source taken in the last 24 h", "DBA", "", ""],
            ["Rollback plan reviewed and rollback owner on call", "Migration lead", "", ""],
            ["Business owner and key users available for validation", "App owner", "", ""],
            ["Change approved and communications sent", "Change manager", "", ""],
            ["No open P1 / P2 incidents on dependent systems", "Ops lead", "", ""],
            ["Monitoring and alerting live on target", "Ops lead", "", ""],
            ["Security sign-off (exposure, access, logging)", "Security lead", "", ""],
        ],
    )
    table_sheet(
        wb,
        "Rollback Plan",
        [("Ref", 8), ("Trigger", 44), ("Rollback action", 56), ("Owner", 16), ("Duration (min)", 12), ("Point of no return?", 14)],
        [
            ["RB-0", "Pre-cutover readiness not met", "Postpone wave; no changes made to production.", "Migration lead", 0, "No"],
            ["RB-1", "Unable to stop source services cleanly", "Restart source services; re-open access; reschedule.", "App support", 15, "No"],
            ["RB-2", "Final sync incomplete or validation counts mismatch", "Set source read-write; resume source; investigate replication.", "DBA", 20, "No"],
            ["RB-3", "Target launch or DB switchover fails", "Abandon target; revert source to read-write; restore DNS.", "Migration engineer", 30, "No"],
            ["RB-4", "Smoke tests or business validation fail and can't be fixed in the window", "Revert DNS / load balancer; re-enable source; capture target state for analysis.", "Network lead", 30, "No"],
            ["RB-5", "Go-live decision is no-go", "Execute RB-4; if writes occurred on target, reverse-sync or replay per DB plan.", "Migration lead", 60, "After go-live + writes: YES"],
        ],
    )
    table_sheet(
        wb,
        "Contacts",
        [("Role", 28), ("Name", 22), ("Phone", 16), ("Email", 28), ("Backup", 22)],
        [["Migration lead", "", "", "", ""], ["Rollback decision owner", "", "", "", ""], ["DBA", "", "", "", ""], ["Network lead", "", "", "", ""], ["Application owner", "", "", "", ""], ["Security lead", "", "", "", ""], [f"{p['short']} support case", "", "", "", ""]],
    )
    fname = f"{slug}-cutover-runbook.xlsx"
    save(wb, os.path.join(OUT, slug, fname))
    return fname


# ---------------------------------------------------------------- commercial checklist

def build_commercial(p: dict, slug: str):
    wb = Workbook()
    cover_sheet(
        wb,
        f"Commercial & Funding Checklist — {p['short']}",
        p,
        "Qualify the opportunity, choose the buying route, plan commitments and secure vendor funding.",
        ["Answer the qualification questions with the customer.", "Compare buying routes on the Commercial Options sheet.", "Track each funding programme from application to evidence.", "Model commitments on the Commitment Ladder after right-sizing."],
    )
    table_sheet(
        wb,
        "Qualification",
        [("Topic", 18), ("Question", 60), ("Answer", 40), ("Impact on deal", 30)],
        [
            ["Driver", "What event forces a decision (DC exit, hardware refresh, licence renewal, audit, new product)?", "", "Timeline and urgency"],
            ["Driver", "What happens if the customer does nothing for 12 months?", "", "Cost of inaction"],
            ["Scope", "How many servers, databases, applications and TB are in scope?", "", "Deal size, funding tier"],
            ["Scope", "Which workloads are regulated (DPDP, RBI, SEBI, IRDAI)?", "", "Region and controls"],
            ["Economics", "Current annual run cost and upcoming capital refresh?", "", "Business case baseline"],
            ["Economics", "Existing licences and support renewal dates?", "", "BYOL and timing"],
            ["Decision", "Who signs, who influences, who can block (security, finance, procurement)?", "", "Stakeholder plan"],
            ["Decision", "Procurement route and approvals needed (tender, GeM, board)?", "", "Close date"],
            ["Competition", "Which other clouds or partners are in the evaluation?", "", "Positioning"],
            ["Funding", f"Is the customer eligible for {p['short']} migration funding or credits?", "", "Net cost to customer"],
        ],
    )
    table_sheet(
        wb,
        "Commercial Options",
        [("Criterion", 32), (f"Direct with {p['short']}", 30), ("Through partner / reseller", 30), ("Marketplace / private offer", 30)],
        [
            ["Billing currency & GST", "", "", ""], ["Consolidated invoice with services", "", "", ""], ["Commitment / discount access", "", "", ""],
            ["Support model", "", "", ""], ["Credit terms", "", "", ""], ["Draw-down against existing commit", "", "", ""], ["Recommendation", "", "", ""],
        ],
    )
    table_sheet(
        wb,
        "Funding Programmes",
        [("Programme", 40), ("What it funds", 36), ("Eligibility", 40), ("Evidence required", 34), ("Status", 13), ("Value (INR)", 14), ("Notes", 30)],
        [[n, w, e, ev, "Not started", "", ""] for (n, w, e, ev) in p["funding"]],
        status_col="Status",
    )
    cl = table_sheet(
        wb,
        "Commitment Ladder",
        [("Input / output", 44), ("Value", 16), ("Notes", 50)],
        [
            ["Steady-state monthly compute on-demand (INR)", 1000000, "After right-sizing, from 30–60 days of telemetry"],
            ["Tranche 1 coverage %", 0.5, "Commit early to the floor of usage"],
            ["Tranche 2 coverage % (after 90 days)", 0.2, "Add once usage is stable"],
            ["Effective discount on committed spend %", 0.3, p["commitments"]],
            ["Monthly saving — tranche 1", None, ""],
            ["Monthly saving — tranche 2", None, ""],
            ["Annual saving (both tranches)", None, ""],
            ["Coverage after both tranches", None, "Keep ≤ 80% to retain flexibility"],
        ],
        freeze="B2",
    )
    cl["B6"] = "=B2*B3*B5"
    cl["B7"] = "=B2*B4*B5"
    cl["B8"] = "=(B6+B7)*12"
    cl["B9"] = "=B3+B4"
    for r in (2, 6, 7, 8):
        cl.cell(row=r, column=2).number_format = "#,##0"
    for r in (3, 4, 5, 9):
        cl.cell(row=r, column=2).number_format = "0%"
    table_sheet(
        wb,
        "Deal Checklist",
        [("Step", 8), ("Item", 56), ("Owner", 18), ("Status", 13), ("Notes", 34)],
        [[i + 1, t, "", "Not started", ""] for i, t in enumerate([
            "Qualification completed and documented", "Discovery data sufficient for pricing", "TCO and business case reviewed with customer finance",
            "Funding applications submitted", "Buying route agreed", "Commitment plan agreed (post right-sizing)", "Security and regulatory requirements captured in the contract",
            "Exit, data-return and audit clauses reviewed", "SOW and acceptance criteria signed", "Customer consent / attestations for funding captured",
        ])],
        status_col="Status",
    )
    fname = f"{slug}-commercial-funding-checklist.xlsx"
    save(wb, os.path.join(OUT, slug, fname))
    return fname


# ---------------------------------------------------------------- docx helpers

def set_cell_bg(cell, hex_color: str):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def new_doc(title: str, subtitle: str, p: dict) -> Document:
    d = Document()
    for s in d.sections:
        s.left_margin = s.right_margin = Cm(2.2)
        s.top_margin = s.bottom_margin = Cm(2)
    st = d.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)
    st.font.color.rgb = RGBColor.from_string(INK)
    for name, size in (("Heading 1", 18), ("Heading 2", 13.5), ("Heading 3", 11.5)):
        h = d.styles[name]
        h.font.name = "Calibri"
        h.font.size = Pt(size)
        h.font.bold = True
        h.font.color.rgb = RGBColor.from_string(INK)
    para = d.add_paragraph()
    r = para.add_run("CloudMigrate.in")
    r.bold = True
    r.font.color.rgb = RGBColor.from_string(ACCENT)
    r.font.size = Pt(11)
    t = d.add_paragraph()
    tr = t.add_run(title)
    tr.bold = True
    tr.font.size = Pt(26)
    s = d.add_paragraph()
    sr = s.add_run(subtitle)
    sr.font.size = Pt(12)
    sr.font.color.rgb = RGBColor.from_string("676C6E")
    meta = d.add_table(rows=4, cols=2)
    meta.style = "Table Grid"
    for i, (k, v) in enumerate([("Cloud", p["name"]), ("Version", f"{VERSION} — {TODAY}"), ("Customer", "[Customer name]"), ("Owner", "[Programme / migration lead]")]):
        meta.cell(i, 0).text = k
        meta.cell(i, 1).text = v
        set_cell_bg(meta.cell(i, 0), "EBE7DC")
    d.add_paragraph()
    note = d.add_paragraph()
    nr = note.add_run("Template provided as-is. Replace bracketed text, delete guidance you do not need, and validate vendor programme terms before use.")
    nr.italic = True
    nr.font.size = Pt(9)
    nr.font.color.rgb = RGBColor.from_string("676C6E")
    return d


def add_table(d: Document, headers: list[str], rows: list[list[str]], widths: list[float] | None = None):
    t = d.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]
        c.text = ""
        run = c.paragraphs[0].add_run(h)
        run.bold = True
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        run.font.size = Pt(9.5)
        set_cell_bg(c, INK)
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = str(v)
            for para in cells[i].paragraphs:
                for run in para.runs:
                    run.font.size = Pt(9.5)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
    d.add_paragraph()
    return t


def bullets(d: Document, items: list[str]):
    for i in items:
        d.add_paragraph(i, style="List Bullet")


# ---------------------------------------------------------------- migration handbook

def build_handbook(p: dict, slug: str):
    d = new_doc("Migration Handbook", f"How we plan, execute and hand over migrations to {p['name']}", p)
    d.add_heading("1. Purpose and scope", level=1)
    d.add_paragraph(
        f"This handbook is the single operating guide for migrating workloads to {p['name']}. It defines the phases, roles, standards, "
        "tools, quality gates and evidence that every wave follows. It is written for programme leads, architects, migration engineers, "
        "DBAs, security and operations teams, and for customer stakeholders who approve each stage."
    )
    d.add_paragraph("In scope: [list applications, environments, data centres]. Out of scope: [list]. Target completion: [date].")
    d.add_heading("2. Principles", level=1)
    bullets(d, [
        "Foundation before workloads — no production workload lands before the landing zone, identity, network, logging and backup baselines are live.",
        "Measure, don't guess — size from at least 14 days of utilisation data (p95), and validate dependencies with application owners.",
        "Same runbook every wave — repeatability beats heroics; durations are updated after every rehearsal.",
        "Rollback is designed, rehearsed and owned — every cutover has a named rollback decision owner and a deadline.",
        "Security and compliance are built in — residency, encryption, logging (≥ 180 days in India) and least privilege from day one.",
        "Evidence as we go — funding, audit and acceptance evidence is captured per wave, not reconstructed later.",
        "Optimise after landing — commitments and right-sizing follow 30–60 days of real telemetry.",
    ])
    d.add_heading("3. Governance and roles", level=1)
    add_table(d, ["Role", "Responsibilities"], [
        ["Executive sponsor", "Owns the business case, resolves escalations, chairs steering committee."],
        ["Programme lead", "Plan, budget, RAID log, wave board, vendor coordination, reporting."],
        ["Lead architect", "Target architecture, landing zone, design authority, technical risk."],
        ["Migration engineers", "Replication, test migrations, cutover execution, runbooks."],
        ["DBA", "Database migration method, validation, performance, rollback for data."],
        ["Security lead", "Controls, approvals, logging, incident response readiness, regulatory mapping."],
        ["Network lead", "Connectivity, IP plan, DNS, firewall rules, load-balancer changes."],
        ["Application owners", "Disposition sign-off, test scripts, business validation, go-live acceptance."],
        ["Operations lead", "Run-team readiness, monitoring, hypercare, handover acceptance."],
    ], [5, 11.5])
    d.add_paragraph("Cadence: daily wave stand-up during execution; weekly design authority; fortnightly steering committee; go / no-go before every cutover.")
    d.add_heading("4. Migration lifecycle", level=1)
    phases = [
        ("4.1 Assess", f"Discovery with {p['discovery']}; application portfolio; dependency mapping; 7R dispositions; TCO and business case; wave plan.", ["Discovery data set", "Application portfolio with owners", "Business case", "Wave plan"]),
        ("4.2 Mobilise / foundation", f"{p['landing_zone']}; {p['identity']}; {p['network_hub']}; {p['connectivity']}; logging, security and backup baselines; pilot wave.", ["Landing zone as code", "Connectivity tested", "Security baseline", "Pilot wave completed"]),
        ("4.3 Migrate (waves)", f"Server replication with {p['replication']}; databases with {p['db_migration']}; bulk data with {p['bulk_transfer']}; VMware with {p['vmware']}.", ["Per-wave runbook", "Test migration results", "Cutover sign-off", "Reconciliation evidence"]),
        ("4.4 Hypercare and handover", "72-hour intensive monitoring, daily owner check-ins, defect triage, run-team handover.", ["Hypercare log", "Handover acceptance", "Updated CMDB"]),
        ("4.5 Optimise", f"Right-sizing, {p['commitments']}, modernisation backlog, decommissioning of source systems.", ["Optimisation report", "Commitment plan", "Decommission certificates"]),
    ]
    for h, body, outs in phases:
        d.add_heading(h, level=2)
        d.add_paragraph(body)
        d.add_paragraph("Exit criteria / outputs:")
        bullets(d, outs)
    d.add_heading("5. Landing-zone reference", level=1)
    add_table(d, ["Capability", f"{p['short']} standard"], [
        ["Regions", p["regions"]], ["Structure & guardrails", p["landing_zone"]], ["Identity", p["identity"]], ["Network hub", p["network_hub"]],
        ["Hybrid connectivity", p["connectivity"]], ["DNS", p["dns"]], ["Logging", p["logging"]], ["Security posture", p["posture"]],
        ["Threat detection", p["threat"]], ["SIEM", p["siem"]], ["Keys & secrets", p["kms"]], ["Backup", p["backup"]], ["DR", p["dr"]],
        ["Patching", p["patching"]], ["Infrastructure as code", p["iac"]], ["Tagging", p["tagging"]], ["Cost management", p["cost"]],
    ], [5, 11.5])
    d.add_heading("6. Migration patterns (7R)", level=1)
    add_table(d, ["Pattern", "When to use", f"Typical {p['short']} approach"], [
        ["Retire", "No users, duplicate capability, or replaced by SaaS", "Archive data per retention policy; decommission"],
        ["Retain", "Hard dependency, licence or latency constraint, or near end-of-life", "Keep on-premises; connect via hybrid network"],
        ["Rehost", "Fast exit, stable apps, low change appetite", p["replication"]],
        ["Relocate", "Large VMware estates with a near-term exit date", p["vmware"]],
        ["Replatform", "Databases and middleware where managed services cut toil", p["db_migration"]],
        ["Repurchase", "Commodity capability better served by SaaS", "Data migration to SaaS; integration rebuild"],
        ["Refactor", "Strategic apps needing scale, agility or AI", "Containers / serverless / managed data services"],
    ], [3, 6, 7.5])
    d.add_heading("7. Wave planning", level=1)
    bullets(d, [
        "Build move groups from validated dependencies; keep chatty components together.",
        "Start with a pilot wave of 3–5 low-risk applications to prove tooling and runbooks.",
        "Typical factory cadence: one wave every two weeks, 10–40 servers per wave once stable.",
        "Respect blackout periods (month-end, quarter-end, festive peaks, audits) and licence / contract dates.",
        "Each wave needs: owner sign-off, test plan, rehearsed cutover, rollback owner, communications.",
    ])
    d.add_heading("8. Cutover procedure", level=1)
    add_table(d, ["When", "Activity"], [
        ["T-7 days", "Scope frozen, change approved, communications sent, test migration complete"],
        ["T-2 days", "DNS TTLs lowered; replication lag verified; backup of source verified"],
        ["T-0 go / no-go", "Readiness checklist reviewed with quorum; rollback owner confirmed"],
        ["Freeze", "Users off; services and jobs stopped; final sync; source read-only"],
        ["Switch", "Target promoted; database switchover; config / secrets updated; DNS / load balancer switched"],
        ["Validate", "Smoke tests; data reconciliation; business validation"],
        ["Decide", "Go-live or rollback before the agreed deadline"],
        ["Go-live", "Access restored; jobs enabled; hypercare starts"],
    ], [3.5, 13])
    d.add_heading("9. Rollback", level=1)
    d.add_paragraph(
        "Rollback triggers and the decision owner are agreed before the cutover window opens. Until the point of no return (target accepts business writes "
        "that cannot be replayed to source), rollback means re-enabling the source and reverting DNS / load balancers. After it, rollback requires a reverse-sync "
        "or data replay plan defined by the DBA. The cutover runbook template contains the standard triggers RB-0 to RB-5."
    )
    d.add_heading("10. Testing and validation", level=1)
    add_table(d, ["Test", "Owner", "Pass criteria"], [
        ["Technical smoke test", "Migration engineer", "Services up, ports reachable, logs and metrics flowing"],
        ["Data reconciliation", "DBA + data owner", "Row counts and checksums match; business control totals agree"],
        ["Functional / UAT", "Application owner", "Signed test scripts for critical journeys"],
        ["Performance", "Architect", "Key transactions within ± 10% of baseline"],
        ["Security", "Security lead", "No unexpected exposure; logging and access verified"],
        ["Backup / restore", "Ops", "First backup completed; restore test passed"],
    ], [4, 4, 8.5])
    d.add_heading("11. Hypercare and handover", level=1)
    bullets(d, [
        "Hypercare: 72 hours intensive, then business-hours support until exit criteria are met (no open P1 / P2, 10 business days stable).",
        "Handover pack: as-built diagrams, runbooks, RACI, SLAs, monitoring dashboards, known-issues log, credentials vaulted.",
        "Operations acceptance is a signed form — not an email.",
    ])
    d.add_heading("12. Top risks and mitigations", level=1)
    add_table(d, ["Risk", "Mitigation"], [
        ["Undiscovered dependencies", "Agent-based dependency mapping + owner validation + extended test migrations"],
        ["Overlapping IP ranges / DNS issues", "IPAM review in foundation phase; hybrid DNS tested both ways"],
        ["Replication bandwidth insufficient", "Size for peaks; schedule initial sync; offline transfer for bulk data"],
        ["Hard-coded IPs / hostnames", "Config scanning; DNS aliases; regression tests"],
        ["Licence non-compliance", "Licence position per workload before migration; BYOL rules checked"],
        ["Performance regression", "Baseline before, measure after; right-size with headroom for first 30 days"],
        ["Data loss or corruption", "Verified backups, validation scripts, read-only source during cutover"],
        ["Security gaps at go-live", "Landing-zone guardrails; security sign-off in go / no-go"],
        ["Regulatory breach (residency / logging)", "Region policy; in-country log retention ≥ 180 days; DPDP review"],
        ["Business disruption", "Blackout calendar; communications; rehearsed rollback"],
    ], [6, 10.5])
    d.add_heading("13. Compliance notes for India", level=1)
    bullets(d, [
        "Digital Personal Data Protection Act 2023 and Rules — reasonable security safeguards, breach notification, purpose limitation, data-principal rights.",
        "CERT-In Directions (April 2022) — report specified incidents within 6 hours; retain ICT logs for 180 days within India; synchronise clocks to approved NTP sources.",
        "Sector regulators — RBI (IT outsourcing, cyber resilience), SEBI (cloud services framework), IRDAI (information and cyber security guidelines) where applicable.",
        "Government workloads — confirm MeitY empanelment of the specific cloud service and region.",
    ])
    d.add_heading("14. Templates referenced", level=1)
    bullets(d, ["Discovery & Assessment Workbook (.xlsx)", "Pre-Migration Checklist (.xlsx)", "Cutover Runbook (.xlsx)", "Post-Migration Checklist (.xlsx)", "Managed Services Runbook, RACI & SLA (.xlsx)", "Security Assessment Checklist (.xlsx)"])
    d.add_heading("Appendix A — Glossary", level=1)
    add_table(d, ["Term", "Meaning"], [
        ["7R", "Retire, Retain, Rehost, Relocate, Replatform, Repurchase, Refactor"], ["CDC", "Change data capture — continuous replication of database changes"],
        ["Move group", "Set of servers / applications migrated together because of dependencies"], ["RTO / RPO", "Recovery time objective / recovery point objective"],
        ["Hypercare", "Heightened support period immediately after go-live"], ["Landing zone", "Pre-configured, governed cloud foundation for workloads"],
    ], [4, 12.5])
    fname = f"{slug}-migration-handbook.docx"
    d.save(os.path.join(OUT, slug, fname))
    return fname


# ---------------------------------------------------------------- POC charter

def build_poc(p: dict, slug: str):
    ex = p["poc_example"]
    d = new_doc("Proof of Concept Charter", f"Template with a worked example for {p['name']}", p)
    d.add_heading("1. Summary", level=1)
    add_table(d, ["Field", "Entry"], [
        ["POC title", ex["title"]], ["Business sponsor", "[Name, role]"], ["Technical lead", "[Name, role]"], ["Start / decision date", "[dd-mmm-yyyy] / [dd-mmm-yyyy]"],
        ["Time box", "4–6 weeks"], ["Environment", f"Isolated sandbox in {p['name']} with budget alerts and guardrails"],
    ], [4.5, 12])
    d.add_heading("2. Business question and hypothesis", level=1)
    d.add_paragraph("Business question: [What decision will this POC enable?]")
    d.add_paragraph(f"Hypothesis (example): {ex['hypothesis']}")
    d.add_heading("3. Success criteria", level=1)
    d.add_paragraph("Each criterion must be measurable and agreed by the sponsor before build starts.")
    add_table(d, ["#", "Criterion", "Target", "Measurement method", "Result"], [[str(i + 1), c, t, "[How it will be measured]", ""] for i, (c, t) in enumerate(ex["criteria"])], [1, 4, 5, 4.5, 2])
    d.add_heading("4. Scope", level=1)
    d.add_paragraph("In scope:")
    bullets(d, ["[Use cases / user journeys]", "[Data sets and volumes]", "[Integrations]"])
    d.add_paragraph("Out of scope:")
    bullets(d, ["Production cut-over", "Full HA / DR (designed, not built)", "[Other exclusions]"])
    d.add_heading("5. Architecture (POC)", level=1)
    d.add_paragraph(f"Services (example): {ex['services']}.")
    d.add_paragraph(f"Guardrails: {p['landing_zone']}; {p['identity']}; budgets and alerts; data classification applied to POC data; no production credentials.")
    d.add_heading("6. Data", level=1)
    add_table(d, ["Data set", "Classification", "Volume", "Source", "Handling"], [["[e.g. policy documents]", "[Internal / Personal]", "[GB / rows]", "[System]", "[Masked / synthetic / approved copy]"]], [3.5, 3, 2.5, 3, 4.5])
    d.add_heading("7. Plan", level=1)
    add_table(d, ["Week", "Activities", "Output"], [
        ["1", "Charter sign-off, environment, data access, test set", "Signed charter, sandbox ready"],
        ["2–3", "Build minimum architecture", "Working POC"],
        ["4", "Testing against criteria; user feedback", "Results log"],
        ["5", "Cost model, production architecture", "Production design and estimate"],
        ["6", "Read-out and decision", "Go / no-go"],
    ], [1.5, 9, 6])
    d.add_heading("8. Roles", level=1)
    add_table(d, ["Role", "Name", "Responsibility"], [["Sponsor", "", "Owns the decision"], ["Technical lead", "", "Design and delivery"], ["Data owner", "", "Approves data use"], ["Security", "", "Approves guardrails"], ["Users / testers", "", "Evaluate outcomes"]], [4, 4, 8.5])
    d.add_heading("9. Risks and assumptions", level=1)
    bullets(d, ["Data access approved by [date].", "Test users available in week 4.", "Vendor quotas / service availability in the target region confirmed.", "[Other]"])
    d.add_heading("10. Funding", level=1)
    funding = p["funding"][:3] if p["funding"] else []
    if funding:
        bullets(d, [f"{n} — {w}" for (n, w, _, _) in funding])
    else:
        d.add_paragraph("Check partner / vendor funding eligibility (see the CloudMigrate Funding Advisor for Microsoft Azure programmes).")
    d.add_heading("11. Decision record", level=1)
    add_table(d, ["Decision", "Rationale", "Next steps", "Signed (sponsor)", "Date"], [["[Go / No-go / Extend]", "", "", "", ""]], [3, 4.5, 4.5, 3, 2])
    d.add_heading("12. Sign-off", level=1)
    add_table(d, ["Role", "Name", "Signature", "Date"], [["Business sponsor", "", "", ""], ["Technical lead", "", "", ""], ["Security", "", "", ""]], [4, 4.5, 5, 3])
    fname = f"{slug}-poc-charter.docx"
    d.save(os.path.join(OUT, slug, fname))
    return fname


# ---------------------------------------------------------------- manifest

META = {
    "handbook": ("Migration Handbook", "docx", "Phases, roles, landing-zone standards, 7R patterns, cutover, rollback, hypercare and India compliance notes."),
    "pre-checklist": ("Pre-Migration Checklist", "xlsx", "Wave readiness gate with owners, priorities, status drop-downs and an auto-calculated summary."),
    "post-checklist": ("Post-Migration Checklist", "xlsx", "Validation, security, backup, cost and handover gate after each wave."),
    "assessment-workbook": ("Discovery & Assessment Workbook", "xlsx", "Server inventory, application portfolio, weighted 7R matrix, dependencies, 3-year TCO and wave plan."),
    "security-checklist": ("Security Assessment Checklist", "xlsx", "Controls mapped to the cloud benchmark, CIS v8 and Indian regulation (DPDP, CERT-In, RBI, SEBI)."),
    "runbook-raci": ("Managed Services Runbook, RACI & SLA", "xlsx", "RACI, SLA matrix, recurring operations runbook, escalation matrix and monthly KPIs."),
    "cutover-runbook": ("Cutover Runbook", "xlsx", "Timed cutover plan with go / no-go gates, rollback triggers and contact sheet."),
    "poc-charter": ("POC Charter", "docx", "Hypothesis, measurable success criteria, scope, plan and decision record — with a worked example."),
    "commercial-checklist": ("Commercial & Funding Checklist", "xlsx", "Qualification questions, buying-route comparison, funding tracker and commitment ladder."),
}


def main():
    manifest: dict = {"generated": dt.date.today().isoformat(), "version": VERSION, "clouds": {}}
    for slug, p in PROFILES.items():
        os.makedirs(os.path.join(OUT, slug), exist_ok=True)
        files = {}
        files["handbook"] = build_handbook(p, slug)
        files["pre-checklist"], _ = build_checklist(p, slug, "pre")
        files["post-checklist"], _ = build_checklist(p, slug, "post")
        files["assessment-workbook"] = build_assessment(p, slug)
        files["security-checklist"], _ = build_security(p, slug)
        files["runbook-raci"] = build_runbook(p, slug)
        files["cutover-runbook"] = build_cutover(p, slug)
        files["poc-charter"] = build_poc(p, slug)
        if slug != "azure":  # Azure sales is served by the Funding Advisor, not a template
            files["commercial-checklist"] = build_commercial(p, slug)
        entries = {}
        for key, fname in files.items():
            title, ftype, desc = META[key]
            size = os.path.getsize(os.path.join(OUT, slug, fname))
            entries[key] = {"title": title, "type": ftype, "description": desc, "href": f"/downloads/{slug}/{fname}", "bytes": size}
        manifest["clouds"][slug] = entries
        print(f"{slug}: {len(files)} files")
    with open(MANIFEST, "w") as f:
        json.dump(manifest, f, indent=2)
    print("manifest ->", MANIFEST)


if __name__ == "__main__":
    main()
